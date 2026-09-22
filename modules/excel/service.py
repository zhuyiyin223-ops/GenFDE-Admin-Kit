"""系统能力：Excel 导出服务。

把行数据写成 xlsx 字节；不查询业务数据，不触发浏览器下载。
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


@dataclass(frozen=True, slots=True)
class ExcelFile:
    """一次导出生成的 Excel 文件。"""

    content: bytes
    filename: str
    row_count: int
    truncated: bool = False


class ExcelExportService:
    """将行数据写成 xlsx。"""

    MAX_ROWS = 65535

    @staticmethod
    def build_filename(stem: str) -> str:
        """按模块名和时间戳生成下载文件名。"""
        safe_stem = str(stem or "export").strip() or "export"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{safe_stem}_{timestamp}.xlsx"

    @classmethod
    def build_file(
        cls,
        rows: list[dict[str, Any]],
        headers: dict[str, str],
        *,
        filename_stem: str,
        total: int | None = None,
        sheet_name: str = "Sheet1",
    ) -> ExcelFile:
        """按表头写出文件；超过上限时截断。"""
        if not headers:
            raise ValueError("导出表头不能为空")
        capped_rows = list(rows)[: cls.MAX_ROWS]
        total_count = len(capped_rows) if total is None else max(0, int(total))
        if not capped_rows:
            return ExcelFile(
                content=b"",
                filename=cls.build_filename(filename_stem),
                row_count=0,
            )
        content = cls.export_rows(capped_rows, headers, sheet_name=sheet_name)
        return ExcelFile(
            content=content,
            filename=cls.build_filename(filename_stem),
            row_count=len(capped_rows),
            truncated=total_count > len(capped_rows),
        )

    @classmethod
    def export_rows(
        cls,
        rows: list[dict[str, Any]],
        headers: dict[str, str],
        *,
        sheet_name: str = "Sheet1",
    ) -> bytes:
        """把行数据写入单个工作表，返回 xlsx 字节。"""
        workbook = Workbook()
        worksheet = workbook.active
        title = str(sheet_name or "Sheet1").strip() or "Sheet1"
        worksheet.title = title[:31]
        cls._write_sheet(worksheet, rows, headers)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _write_sheet(
        worksheet: Worksheet,
        rows: list[dict[str, Any]],
        headers: dict[str, str],
    ) -> None:
        """写入表头与数据行。"""
        header_keys = list(headers.keys())
        header_font = Font(bold=True)
        for column_index, key in enumerate(header_keys, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=headers[key])
            cell.font = header_font

        for row_index, row in enumerate(rows, start=2):
            for column_index, key in enumerate(header_keys, start=1):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=_cell_value(row.get(key)),
                )

        worksheet.freeze_panes = "A2"
        for column_index, key in enumerate(header_keys, start=1):
            sample_lengths = [len(str(headers[key]))]
            for row in rows[:50]:
                sample_lengths.append(len(str(_cell_value(row.get(key)))))
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(sample_lengths) + 2,
                40,
            )


def _cell_value(value: Any) -> Any:
    """把单元格值转成 openpyxl 可写入的标量。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (list, tuple, dict, set)):
        return "、".join(str(item) for item in value) if not isinstance(value, dict) else str(value)
    return value
